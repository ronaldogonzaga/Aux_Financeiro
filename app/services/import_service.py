from __future__ import annotations
import csv
from io import BytesIO, StringIO
from typing import Any
from odf.opendocument import load as load_ods
from odf.table import Table, TableCell, TableRow
from odf.text import P
from openpyxl import load_workbook
COLUNAS_NOME = {'nome', 'name', 'medico', 'médico', 'prestador', 'doctor'}
COLUNAS_EMAIL = {'email', 'e-mail', 'mail', 'correo'}

def _normalize_header(value: Any) -> str:
    if value is None:
        return ''
    return str(value).strip().lower()

def _map_columns(headers: list[str]) -> tuple[int | None, int | None]:
    idx_nome = None
    idx_email = None
    for i, h in enumerate(headers):
        if h in COLUNAS_NOME and idx_nome is None:
            idx_nome = i
        if h in COLUNAS_EMAIL and idx_email is None:
            idx_email = i
    if idx_nome is None and len(headers) >= 1:
        idx_nome = 0
    if idx_email is None and len(headers) >= 2:
        idx_email = 1
    return (idx_nome, idx_email)

def _itens_from_rows(rows: list[list[Any]]) -> list[dict[str, str]]:
    if not rows:
        return []
    headers = [_normalize_header(c) for c in rows[0]]
    idx_nome, idx_email = _map_columns(headers)
    if idx_nome is None or idx_email is None:
        raise ValueError('Planilha deve conter colunas de Nome e E-mail.')
    itens: list[dict[str, str]] = []
    for row in rows[1:]:
        if not row:
            continue
        nome = str(row[idx_nome] or '').strip() if idx_nome < len(row) else ''
        email = str(row[idx_email] or '').strip() if idx_email < len(row) else ''
        if nome and email:
            itens.append({'nome': nome, 'email': email})
    return itens

def parse_xlsx(content: bytes) -> list[dict[str, str]]:
    wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return _itens_from_rows(rows)

def _ods_cell_text(cell: TableCell) -> str:
    partes = []
    for p in cell.getElementsByType(P):
        partes.append(str(p))
    return ''.join(partes).strip()

def parse_ods(content: bytes) -> list[dict[str, str]]:
    doc = load_ods(BytesIO(content))
    tables = doc.spreadsheet.getElementsByType(Table)
    if not tables:
        return []
    sheet = tables[0]
    rows: list[list[Any]] = []
    for row in sheet.getElementsByType(TableRow):
        cells: list[Any] = []
        for cell in row.getElementsByType(TableCell):
            repeat = int(cell.getAttribute('numbercolumnsrepeated') or 1)
            repeat = min(repeat, 50)
            texto = _ods_cell_text(cell)
            for _ in range(repeat):
                cells.append(texto)
        while cells and (not cells[-1]):
            cells.pop()
        if cells:
            rows.append(cells)
    return _itens_from_rows(rows)

def parse_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode('utf-8-sig', errors='replace')
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=';,\t')
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ';' if sample.count(';') >= sample.count(',') else ','
    reader = csv.reader(StringIO(text), dialect)
    rows = list(reader)
    return _itens_from_rows(rows)

def parse_planilha(filename: str, content: bytes) -> list[dict[str, str]]:
    lower = filename.lower()
    if lower.endswith(('.xlsx', '.xlsm')):
        return parse_xlsx(content)
    if lower.endswith('.ods'):
        return parse_ods(content)
    if lower.endswith('.csv'):
        return parse_csv(content)
    raise ValueError('Formato não suportado. Use .xlsx, .ods ou .csv')
